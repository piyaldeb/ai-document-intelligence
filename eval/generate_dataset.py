"""Generate the labeled evaluation dataset (deterministic, seeded).

Produces realistic business documents with known ground truth:
  * digital PDFs in 5 layout templates (label synonyms, 5 currencies, US/EU number and date
    formats, multi-page invoices with continued tables),
  * a multi-invoice batch PDF (segmentation),
  * scanned PDFs and JPG/PNG/TIFF photos with degradations (noise, blur, skew, rotation,
    low resolution, low contrast, JPEG artefacts),
  * XLSX (multi-sheet register with title rows, line-level export), XLS, CSV (semicolon +
    decimal comma + latin-1, and comma + UTF-8),
  * corrupted / encrypted / empty files,
  * validation scenarios (arithmetic errors, duplicates, missing fields),
  * a PDF-vs-Excel cross-verification pair with planted discrepancies,
  * optional stress files (1000-page digital PDF, 100-page scanned PDF) with --stress.

Usage (inside the backend container, which has reportlab/PyMuPDF/Pillow/openpyxl/xlwt):
    python eval/generate_dataset.py --out eval/dataset [--stress]
"""

from __future__ import annotations

import argparse
import copy
import csv
import io
import json
import random
from datetime import date, timedelta
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import numpy as np
import openpyxl
import pymupdf as fitz
import xlwt
from PIL import Image, ImageEnhance, ImageFilter
from reportlab.lib.pagesizes import A4, LETTER
from reportlab.pdfgen import canvas

R = random.Random(20261027)
Q = Decimal("0.01")

SUPPLIERS = [
    "Acme Industrial Supplies Ltd", "Northwind Traders Inc", "Bluewave Logistics LLC", "Greenfield Office Solutions",
    "Delta Textiles Pvt Ltd", "Brightstar Electronics GmbH", "Harbor Freight Services Co.", "Summit Consulting Group",
    "Pioneer Packaging Ltd", "Orion Software Systems", "Meridian Medical Supplies", "Atlas Engineering Works",
    "Crescent Foods Limited", "Silverline Printing Press", "Evergreen Facilities Management", "Kestrel Data Labs",
]
CUSTOMERS = [
    "Riverside Home Care", "Lakeside Manufacturing PLC", "Sunrise Retail Group", "Metro Health Clinic",
    "Oakwood Property Partners", "Vertex Hospitality Ltd", "Hillcrest School Trust", "Cedar Logistics Hub",
    "Nova Pharmaceuticals", "Granite Construction Co", "Bayview Dental Practice", "Unity Microfinance",
]
ITEMS = [
    ("Consulting services", 60, 180), ("Software license (annual)", 200, 1500), ("Steel bolts M8 (box of 100)", 8, 25),
    ("A4 copier paper (ream)", 4, 9), ("Network switch 24-port", 180, 450), ("Delivery and handling", 15, 90),
    ("Cotton fabric roll", 40, 160), ("Cloud hosting (monthly)", 50, 600), ("Safety gloves (pair)", 2, 7),
    ("Printer toner cartridge", 35, 120), ("On-site support visit", 90, 300), ("Office chair ergonomic", 120, 380),
    ("Packaging cartons large", 1, 4), ("Data backup service", 30, 220), ("Medical gauze pack", 3, 12),
    ("Training workshop", 250, 1200), ("LED panel light 60x60", 18, 55), ("Maintenance contract", 150, 900),
]
CURRENCIES = {
    "USD": {"symbol": "$", "tax": [Decimal("0.07"), Decimal("0.0825"), Decimal("0")], "dayfirst": False},
    "EUR": {"symbol": "€", "tax": [Decimal("0.19"), Decimal("0.21"), Decimal("0.20")], "dayfirst": True},
    "GBP": {"symbol": "£", "tax": [Decimal("0.20")], "dayfirst": True},
    "INR": {"symbol": "Rs.", "tax": [Decimal("0.18"), Decimal("0.12")], "dayfirst": True},
    "BDT": {"symbol": "Tk", "tax": [Decimal("0.15"), Decimal("0.075")], "dayfirst": True},
}


def q(x) -> Decimal:
    return Decimal(x).quantize(Q, rounding=ROUND_HALF_UP)


# ------------------------------------------------------------------ ground truth


def make_invoice(n_items: int | None = None, currency: str | None = None, number: str | None = None,
                 supplier: str | None = None, with_tax: bool = True, with_customer: bool = True) -> dict:
    currency = currency or R.choice(list(CURRENCIES))
    n_items = n_items or R.randint(1, 8)
    items = []
    for desc, lo, hi in R.sample(ITEMS, min(n_items, len(ITEMS))):
        qty = Decimal(R.choice([1, 1, 2, 3, 4, 5, 10, 12, 20, 25]))
        unit = q(R.uniform(lo, hi))
        items.append({"description": desc, "quantity": qty, "unit_price": unit, "amount": q(qty * unit)})
    while len(items) < n_items:  # long invoices repeat catalogue items with a suffix
        desc, lo, hi = R.choice(ITEMS)
        qty = Decimal(R.randint(1, 9))
        unit = q(R.uniform(lo, hi))
        items.append({"description": f"{desc} - batch {len(items) + 1}", "quantity": qty, "unit_price": unit,
                      "amount": q(qty * unit)})
    subtotal = q(sum(i["amount"] for i in items))
    rate = R.choice(CURRENCIES[currency]["tax"]) if with_tax else None
    tax = q(subtotal * rate) if rate is not None else None
    total = q(subtotal + (tax or 0))
    d = date(2025, 1, 1) + timedelta(days=R.randint(0, 600))
    return {
        "invoice_number": number or R.choice(["INV-", "INV", "BL-", "", "SI/2025/", "TX-"]) + str(R.randint(1000, 99999)),
        "invoice_date": d,
        "supplier_name": supplier or R.choice(SUPPLIERS),
        "customer_name": R.choice(CUSTOMERS) if with_customer else None,
        "subtotal": subtotal if with_tax else None,
        "tax": tax,
        "tax_rate": rate,
        "total": total,
        "currency": currency,
        "line_items": items,
    }


def gt_json(inv: dict) -> dict:
    def s(v):
        return None if v is None else str(v)

    return {
        "invoice_number": inv["invoice_number"],
        "invoice_date": inv["invoice_date"].isoformat() if inv["invoice_date"] else None,
        "supplier_name": inv["supplier_name"],
        "customer_name": inv["customer_name"],
        "subtotal": s(inv["subtotal"]),
        "tax": s(inv["tax"]),
        "total": s(inv["total"]),
        "currency": inv["currency"],
        "line_items": [{k: s(v) if k != "description" else v for k, v in li.items()} for li in inv["line_items"]],
    }


# ------------------------------------------------------------------ formatting


def fmt_amount(v: Decimal, eu: bool) -> str:
    s = f"{v:,.2f}"
    if eu:
        s = s.replace(",", "X").replace(".", ",").replace("X", ".")
    return s


def fmt_date(d: date, style: str) -> str:
    return {
        "us": d.strftime("%m/%d/%Y"),
        "uk": d.strftime("%d/%m/%Y"),
        "iso": d.isoformat(),
        "long": d.strftime("%B %d, %Y").replace(" 0", " "),
        "short": d.strftime("%d %b %Y"),
        "dots": d.strftime("%d.%m.%Y"),
    }[style]


def date_style_for(currency: str) -> str:
    if currency == "USD":
        return R.choice(["us", "long", "iso"])
    if currency == "EUR":
        return R.choice(["dots", "iso", "short"])
    return R.choice(["uk", "short", "long", "iso"])


# ------------------------------------------------------------------ PDF templates


class Doc:
    def __init__(self, path: Path, pagesize=A4, encrypt: str | None = None):
        self.c = canvas.Canvas(str(path), pagesize=pagesize, encrypt=encrypt)
        self.w, self.h = pagesize

    def text(self, x, y, s, size=10, bold=False, right=False):
        self.c.setFont("Helvetica-Bold" if bold else "Helvetica", size)
        (self.c.drawRightString if right else self.c.drawString)(x, y, s)


def _items_table(d: Doc, inv: dict, y: float, cols: list[tuple[str, float, str]], eu: bool, sym: str,
                 page_info: dict, header_fn, indexed: bool = False) -> float:
    """Draw line items with pagination. Returns the y after the table."""
    def header(yy):
        for label, x, align in cols:
            d.text(x, yy, label, 10, bold=True, right=(align == "r"))
        d.c.line(40, yy - 4, d.w - 40, yy - 4)
        return yy - 18

    y = header(y)
    for idx, li in enumerate(inv["line_items"], start=1):
        if y < 150:
            d.text(d.w - 40, 40, f"Page {page_info['page']} of {page_info['pages']}", 8, right=True)
            d.text(40, 40, "continued on next page", 8)
            d.c.showPage()
            page_info["page"] += 1
            y = header_fn(d) - 10
            y = header(y)
        values = {
            "idx": str(idx), "desc": li["description"], "qty": str(li["quantity"]),
            "unit": fmt_amount(li["unit_price"], eu), "amount": fmt_amount(li["amount"], eu),
        }
        for (label, x, align), key in zip(cols, (["idx"] if indexed else []) + ["desc", "qty", "unit", "amount"]):
            d.text(x, y, values[key], 9.5, right=(align == "r"))
        y -= 16
    d.c.line(40, y + 8, d.w - 40, y + 8)
    return y - 6


def _count_pages(n_items: int, first_capacity: int = 24, next_capacity: int = 34) -> int:
    if n_items <= first_capacity:
        return 1
    return 1 + -(-(n_items - first_capacity) // next_capacity)


def template_classic(path: Path, inv: dict, *, encrypt: str | None = None, extra_page: bool = False) -> None:
    """Letterhead left, INVOICE right, 'Invoice No:' / 'Date:', 'Bill To:', US-style table."""
    cur = inv["currency"]
    eu, sym = False, CURRENCIES[cur]["symbol"]
    d = Doc(path, LETTER if cur == "USD" else A4, encrypt)
    pages = _count_pages(len(inv["line_items"]))
    info = {"page": 1, "pages": pages}

    def head(doc: Doc) -> float:
        doc.text(40, doc.h - 60, inv["supplier_name"], 16, bold=True)
        doc.text(40, doc.h - 76, "12 Commerce Street, Suite 400", 9)
        doc.text(40, doc.h - 88, "accounts@example.com  |  +1 555 0100", 9)
        doc.text(doc.w - 40, doc.h - 60, "INVOICE", 20, bold=True, right=True)
        if inv["invoice_number"]:
            doc.text(doc.w - 200, doc.h - 90, "Invoice No:", 10, bold=True)
            doc.text(doc.w - 40, doc.h - 90, inv["invoice_number"], 10, right=True)
        doc.text(doc.w - 200, doc.h - 104, "Date:", 10, bold=True)
        doc.text(doc.w - 40, doc.h - 104, fmt_date(inv["invoice_date"], inv["_date_style"]), 10, right=True)
        doc.text(doc.w - 200, doc.h - 118, "Due Date:", 10, bold=True)
        doc.text(doc.w - 40, doc.h - 118, fmt_date(inv["invoice_date"] + timedelta(days=30), inv["_date_style"]), 10,
                 right=True)
        return doc.h - 150

    y = head(d)
    if inv["customer_name"]:
        d.text(40, y, "Bill To:", 10, bold=True)
        d.text(40, y - 14, inv["customer_name"], 10)
        d.text(40, y - 28, "45 Market Road, Springfield", 9)
        y -= 60
    cols = [("Description", 40, "l"), ("Qty", 340, "r"), ("Unit Price", 440, "r"), ("Amount", d.w - 40, "r")]
    y = _items_table(d, inv, y, cols, eu, sym, info, head)
    lx, rx = d.w - 220, d.w - 40
    if inv["subtotal"] is not None:
        d.text(lx, y, "Subtotal", 10)
        d.text(rx, y, f"{sym}{fmt_amount(inv['subtotal'], eu)}", 10, right=True)
        y -= 16
    if inv["tax"] is not None:
        pct = f"{(inv['tax_rate'] * 100).normalize():f}%"
        d.text(lx, y, f"Tax ({pct})", 10)
        d.text(rx, y, f"{sym}{fmt_amount(inv['tax'], eu)}", 10, right=True)
        y -= 16
    d.text(lx, y, f"Total ({cur})", 11, bold=True)
    d.text(rx, y, f"{sym}{fmt_amount(inv['total'], eu)}", 11, bold=True, right=True)
    d.text(40, 80, "Payment terms: 30 days. Thank you for your business.", 9)
    if pages > 1:
        d.text(d.w - 40, 40, f"Page {info['page']} of {pages}", 8, right=True)
    d.c.showPage()
    if extra_page:
        d.text(40, d.h - 60, "Terms and Conditions", 12, bold=True)
        for i in range(8):
            d.text(40, d.h - 90 - i * 14, f"{i + 1}. Goods remain the property of the seller until paid in full.", 9)
        d.c.showPage()
    d.c.save()


def template_modern(path: Path, inv: dict, **_) -> None:
    """TAX INVOICE title, From / Bill To side by side, 'Invoice #', 'Issue Date', indexed table, 'Amount Due'."""
    cur = inv["currency"]
    eu, sym = False, cur
    d = Doc(path, A4)
    pages = _count_pages(len(inv["line_items"]), 22)
    info = {"page": 1, "pages": pages}

    def head(doc: Doc) -> float:
        doc.c.setFillColorRGB(0.1, 0.21, 0.36)
        doc.c.rect(0, doc.h - 70, doc.w, 70, fill=1, stroke=0)
        doc.c.setFillColorRGB(1, 1, 1)
        doc.text(40, doc.h - 45, "TAX INVOICE", 22, bold=True)
        doc.c.setFillColorRGB(0, 0, 0)
        doc.text(40, doc.h - 100, "Invoice #", 10, bold=True)
        doc.text(130, doc.h - 100, inv["invoice_number"], 10)
        doc.text(330, doc.h - 100, "Issue Date", 10, bold=True)
        doc.text(420, doc.h - 100, fmt_date(inv["invoice_date"], inv["_date_style"]), 10)
        return doc.h - 130

    y = head(d)
    d.text(40, y, "From", 10, bold=True)
    d.text(330, y, "Bill To", 10, bold=True)
    d.text(40, y - 14, inv["supplier_name"], 10)
    if inv["customer_name"]:
        d.text(330, y - 14, inv["customer_name"], 10)
    d.text(40, y - 28, "Unit 7, Riverside Business Park", 9)
    d.text(330, y - 28, "Accounts Payable Department", 9)
    y -= 60
    cols = [("#", 40, "l"), ("Item", 65, "l"), ("Hours", 330, "r"), ("Rate", 430, "r"), ("Line Total", d.w - 40, "r")]
    y = _items_table(d, inv, y, cols, eu, sym, info, head, indexed=True)
    lx, rx = d.w - 230, d.w - 40
    if inv["subtotal"] is not None:
        d.text(lx, y, "Sub Total", 10)
        d.text(rx, y, f"{fmt_amount(inv['subtotal'], eu)}", 10, right=True)
        y -= 16
    if inv["tax"] is not None:
        d.text(lx, y, f"VAT {(inv['tax_rate'] * 100).normalize():f}%", 10)
        d.text(rx, y, f"{fmt_amount(inv['tax'], eu)}", 10, right=True)
        y -= 16
    d.text(lx, y, "Amount Due", 11, bold=True)
    d.text(rx, y, f"{cur} {fmt_amount(inv['total'], eu)}", 11, bold=True, right=True)
    if pages > 1:
        d.text(d.w - 40, 40, f"Page {info['page']} of {pages}", 8, right=True)
    d.c.showPage()
    d.c.save()


def template_european(path: Path, inv: dict, **_) -> None:
    """EU formatting: 1.234,56, dd.mm.yyyy, 'Invoice number', 'Net amount', 'Total amount EUR'."""
    cur = inv["currency"]
    eu = True
    d = Doc(path, A4)
    info = {"page": 1, "pages": _count_pages(len(inv["line_items"]))}

    def head(doc: Doc) -> float:
        doc.text(doc.w - 40, doc.h - 50, inv["supplier_name"], 13, bold=True, right=True)
        doc.text(doc.w - 40, doc.h - 64, "Hauptstrasse 21, 10115 Berlin", 9, right=True)
        doc.text(doc.w - 40, doc.h - 76, "VAT Reg No: DE123456789", 9, right=True)
        doc.text(40, doc.h - 130, "Invoice", 18, bold=True)
        doc.text(40, doc.h - 155, "Invoice number:", 10)
        doc.text(150, doc.h - 155, inv["invoice_number"], 10)
        doc.text(40, doc.h - 169, "Invoice date:", 10)
        doc.text(150, doc.h - 169, fmt_date(inv["invoice_date"], inv["_date_style"]), 10)
        return doc.h - 200

    y = head(d)
    if inv["customer_name"]:
        d.text(40, d.h - 60, "Customer:", 9, bold=True)
        d.text(40, d.h - 74, inv["customer_name"], 10)
    cols = [("Description", 40, "l"), ("Quantity", 330, "r"), ("Unit price", 430, "r"), ("Amount", d.w - 40, "r")]
    y = _items_table(d, inv, y, cols, eu, "€", info, head)
    lx, rx = d.w - 230, d.w - 40
    if inv["subtotal"] is not None:
        d.text(lx, y, "Net amount", 10)
        d.text(rx, y, f"{fmt_amount(inv['subtotal'], eu)} {cur}", 10, right=True)
        y -= 16
    if inv["tax"] is not None:
        d.text(lx, y, f"VAT {(inv['tax_rate'] * 100).normalize():f} %", 10)
        d.text(rx, y, f"{fmt_amount(inv['tax'], eu)} {cur}", 10, right=True)
        y -= 16
    d.text(lx, y, f"Total amount {cur}", 11, bold=True)
    d.text(rx, y, f"{fmt_amount(inv['total'], eu)}", 11, bold=True, right=True)
    if info["pages"] > 1:
        d.text(d.w - 40, 40, f"Page {info['page']} of {info['pages']}", 8, right=True)
    d.c.showPage()
    d.c.save()


def template_receipt(path: Path, inv: dict, **_) -> None:
    """Narrow till receipt: store name, 'Receipt No.', items, TOTAL; often no customer/tax."""
    cur = inv["currency"]
    sym = CURRENCIES[cur]["symbol"]
    height = 330 + 30 * len(inv["line_items"])
    d = Doc(path, (260, height))
    y = height - 40
    d.text(130 - 5 * len(inv["supplier_name"]) / 2, y, inv["supplier_name"], 10, bold=True)
    y -= 16
    d.text(20, y, "Main Branch, Dhaka Road", 8)
    y -= 24
    d.text(20, y, f"Receipt No.  {inv['invoice_number']}", 9)
    y -= 13
    d.text(20, y, f"Date  {fmt_date(inv['invoice_date'], inv['_date_style'])}", 9)
    y -= 13
    if inv["customer_name"]:
        d.text(20, y, f"Customer  {inv['customer_name']}", 9)
        y -= 13
    y -= 8
    d.text(20, y, "Item", 9, bold=True)
    d.text(240, y, "Amount", 9, bold=True, right=True)
    y -= 14
    for li in inv["line_items"]:
        d.text(20, y, li["description"][:30], 8.5)
        y -= 11
        d.text(30, y, f"{li['quantity']} x {fmt_amount(li['unit_price'], False)}", 8.5)
        d.text(240, y, fmt_amount(li["amount"], False), 8.5, right=True)
        y -= 15
    d.c.line(20, y + 6, 240, y + 6)
    y -= 8
    if inv["subtotal"] is not None:
        d.text(20, y, "Subtotal", 9)
        d.text(240, y, fmt_amount(inv["subtotal"], False), 9, right=True)
        y -= 13
    if inv["tax"] is not None:
        d.text(20, y, f"VAT {(inv['tax_rate'] * 100).normalize():f}%", 9)
        d.text(240, y, fmt_amount(inv["tax"], False), 9, right=True)
        y -= 13
    d.text(20, y, "TOTAL", 11, bold=True)
    d.text(240, y, f"{sym} {fmt_amount(inv['total'], False)}", 11, bold=True, right=True)
    y -= 26
    d.text(20, y, "Thank you for shopping with us!", 8)
    d.c.showPage()
    d.c.save()


def template_statement(path: Path, inv: dict, **_) -> None:
    """Bill-style layout with 'Bill No.', 'Billing Date', 'Customer Name', 'Grand Total'."""
    cur = inv["currency"]
    sym = CURRENCIES[cur]["symbol"]
    d = Doc(path, A4)
    info = {"page": 1, "pages": _count_pages(len(inv["line_items"]))}

    def head(doc: Doc) -> float:
        doc.text(doc.w / 2, doc.h - 50, inv["supplier_name"].upper(), 15, bold=True, right=False)
        doc.text(40, doc.h - 90, "BILL", 14, bold=True)
        doc.text(40, doc.h - 115, f"Bill No.: {inv['invoice_number']}", 10)
        doc.text(330, doc.h - 115, f"Billing Date: {fmt_date(inv['invoice_date'], inv['_date_style'])}", 10)
        return doc.h - 150

    y = head(d)
    if inv["customer_name"]:
        d.text(40, y, f"Customer Name: {inv['customer_name']}", 10)
        y -= 30
    cols = [("Particulars", 40, "l"), ("Qty", 330, "r"), ("Rate", 430, "r"), ("Total", d.w - 40, "r")]
    y = _items_table(d, inv, y, cols, False, sym, info, head)
    lx, rx = d.w - 230, d.w - 40
    if inv["subtotal"] is not None:
        d.text(lx, y, "Sub-total", 10)
        d.text(rx, y, fmt_amount(inv["subtotal"], False), 10, right=True)
        y -= 16
    if inv["tax"] is not None:
        d.text(lx, y, "GST", 10)
        d.text(rx, y, fmt_amount(inv["tax"], False), 10, right=True)
        y -= 16
    d.text(lx, y, "Grand Total", 11, bold=True)
    d.text(rx, y, f"{cur} {fmt_amount(inv['total'], False)}", 11, bold=True, right=True)
    if info["pages"] > 1:
        d.text(d.w - 40, 40, f"Page {info['page']} of {info['pages']}", 8, right=True)
    d.c.showPage()
    d.c.save()


TEMPLATES = {
    "classic": template_classic,
    "modern": template_modern,
    "european": template_european,
    "receipt": template_receipt,
    "statement": template_statement,
}
TEMPLATE_CURRENCIES = {
    "classic": ["USD", "GBP", "USD", "INR"],
    "modern": ["GBP", "USD", "EUR", "BDT"],
    "european": ["EUR"],
    "receipt": ["BDT", "INR", "USD"],
    "statement": ["INR", "BDT", "GBP"],
}


def render(template: str, path: Path, inv: dict, **kw) -> None:
    inv.setdefault("_date_style", date_style_for(inv["currency"]))
    TEMPLATES[template](path, inv, **kw)


# ------------------------------------------------------------------ degradations


def pdf_to_images(pdf: Path, dpi: int) -> list[Image.Image]:
    doc = fitz.open(pdf)
    out = []
    for page in doc:
        pix = page.get_pixmap(dpi=dpi)
        out.append(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
    return out


def degrade(img: Image.Image, level: str, rnd: random.Random) -> Image.Image:
    img = img.convert("L")
    if level in ("scan", "noisy", "skewed", "lowres", "faded"):
        angle = rnd.uniform(-2.5, 2.5) if level != "skewed" else rnd.choice([-4.0, 3.5, 4.5])
        img = img.rotate(angle, resample=Image.BICUBIC, expand=True, fillcolor=255)
    if level in ("noisy", "scan"):
        arr = np.asarray(img).astype(np.int16)
        sigma = 18 if level == "noisy" else 8
        arr = arr + np.random.default_rng(rnd.randint(0, 10**6)).normal(0, sigma, arr.shape).astype(np.int16)
        img = Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8))
    if level == "noisy":
        img = img.filter(ImageFilter.GaussianBlur(0.8))
    if level == "faded":
        img = ImageEnhance.Contrast(img).enhance(0.35)
        img = ImageEnhance.Brightness(img).enhance(1.25)
    if level == "lowres":
        w, h = img.size
        img = img.resize((int(w * 0.45), int(h * 0.45)), Image.BILINEAR)
    if level == "rotated90":
        img = img.rotate(90, expand=True, fillcolor=255)
    return img


def images_to_pdf(images: list[Image.Image], path: Path, dpi: int) -> None:
    first, rest = images[0].convert("RGB"), [i.convert("RGB") for i in images[1:]]
    first.save(path, "PDF", resolution=dpi, save_all=True, append_images=rest)


# ------------------------------------------------------------------ main


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="eval/dataset")
    ap.add_argument("--stress", action="store_true")
    args = ap.parse_args()
    out = Path(args.out)
    for sub in ("pdf", "scanned", "images", "spreadsheets", "corrupted", "validation", "crossverify"):
        (out / sub).mkdir(parents=True, exist_ok=True)
    labels: list[dict] = []
    tmp = out / "_tmp.pdf"

    def add(file: Path, kind: str, invoices: list[dict], *, expect="ok", error_code=None, issues=None, tags=None,
            phase="main"):
        labels.append({
            "file": file.relative_to(out).as_posix(), "kind": kind, "expect": expect, "error_code": error_code,
            "invoices": [gt_json(i) for i in invoices], "expected_issues": issues or [], "tags": tags or [],
            "phase": phase,
        })

    # 1) digital PDFs ------------------------------------------------------------
    n = 0
    for template, currencies in TEMPLATE_CURRENCIES.items():
        for k in range(6):
            cur = currencies[k % len(currencies)]
            with_tax = not (template == "receipt" and k % 2 == 0)
            with_customer = not (template in ("receipt",) and k % 3 != 1) and not (k == 5)
            n_items = R.choice([1, 3, 5, 8]) if k != 4 else R.randint(28, 45)  # one long multi-page invoice each
            if template == "receipt":
                n_items = min(n_items, 6)
            inv = make_invoice(n_items, cur, with_tax=with_tax, with_customer=with_customer)
            n += 1
            f = out / "pdf" / f"{template}_{k + 1:02d}.pdf"
            render(template, f, inv, extra_page=(template == "classic" and k == 2))
            tags = [template, cur] + (["multipage"] if n_items > 24 else []) + ([] if with_tax else ["no_tax"]) + (
                [] if with_customer else ["no_customer"])
            add(f, "pdf", [inv], tags=tags)

    # 2) multi-invoice batch PDF (segmentation) + cross-verification ERP export ---------
    batch = []
    merged = fitz.open()
    for i in range(20):
        template = ["classic", "modern", "statement", "european"][i % 4]
        cur = TEMPLATE_CURRENCIES[template][i % len(TEMPLATE_CURRENCIES[template])]
        inv = make_invoice(R.choice([2, 4, 6]) if i != 9 else 30, cur, number=f"MAR-{2600 + i}")
        render(template, tmp, inv)
        merged.insert_pdf(fitz.open(tmp))
        batch.append(inv)
    batch_file = out / "pdf" / "batch_march_20_invoices.pdf"
    merged.save(batch_file)
    add(batch_file, "pdf", batch, tags=["multi_invoice", "segmentation"])

    erp = [copy.deepcopy(b) for b in batch[:18]]
    erp[2]["total"] = q(erp[2]["total"] + Decimal("50.00"))
    erp[6]["invoice_date"] = erp[6]["invoice_date"] + timedelta(days=1)
    erp[10]["tax"] = q((erp[10]["tax"] or 0) + Decimal("12.50"))
    erp[10]["total"] = q(erp[10]["total"] + Decimal("12.50"))
    erp[4]["supplier_name"] = erp[4]["supplier_name"].replace(" Ltd", " Limited")
    extra = make_invoice(3, "USD", number="MAR-2999")
    erp.append(extra)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AP March"
    ws.append(["Invoice No", "Invoice Date", "Supplier", "Customer", "Net Amount", "VAT", "Total", "Currency"])
    for e in erp:
        ws.append([e["invoice_number"], e["invoice_date"], e["supplier_name"], e["customer_name"],
                   float(e["subtotal"]) if e["subtotal"] is not None else None,
                   float(e["tax"]) if e["tax"] is not None else None, float(e["total"]), e["currency"]])
    erp_file = out / "crossverify" / "erp_export_march.xlsx"
    wb.save(erp_file)
    for e in erp:
        e["line_items"] = []
    add(erp_file, "xlsx", erp, tags=["crossverify"], phase="crossverify")
    expected_recon = {"left": batch_file.relative_to(out).as_posix(), "right": erp_file.relative_to(out).as_posix(),
                      "matched": 15, "mismatch": 3, "missing_in_right": 2, "missing_in_left": 1,
                      "mismatched_invoices": {batch[2]["invoice_number"]: ["total"],
                                              batch[6]["invoice_number"]: ["invoice_date"],
                                              batch[10]["invoice_number"]: ["tax", "total"]}}

    # 3) scanned PDFs -----------------------------------------------------------------
    for i, level in enumerate(["scan", "scan", "scan", "noisy", "skewed", "faded", "scan", "lowres", "noisy", "scan"]):
        template = ["classic", "modern", "statement", "european", "classic"][i % 5]
        cur = TEMPLATE_CURRENCIES[template][i % len(TEMPLATE_CURRENCIES[template])]
        inv = make_invoice(R.choice([2, 3, 5]) if i != 6 else 30, cur)
        render(template, tmp, inv)
        imgs = [degrade(im, level, R) for im in pdf_to_images(tmp, 200)]
        f = out / "scanned" / f"scanned_{template}_{level}_{i + 1:02d}.pdf"
        images_to_pdf(imgs, f, 200)
        add(f, "pdf", [inv], tags=["scanned", level, template])

    # 4) images: JPG / PNG / TIFF -----------------------------------------------------------
    image_plan = [
        ("png", "clean"), ("png", "scan"), ("jpg", "scan"), ("jpg", "noisy"), ("jpg", "lowres"), ("png", "faded"),
        ("jpg", "skewed"), ("png", "rotated90"), ("tiff", "scan"), ("jpg", "clean"), ("png", "noisy"), ("jpg", "scan"),
    ]
    for i, (ext, level) in enumerate(image_plan):
        template = ["receipt", "classic", "modern", "statement"][i % 4]
        cur = TEMPLATE_CURRENCIES[template][i % len(TEMPLATE_CURRENCIES[template])]
        inv = make_invoice(R.choice([1, 2, 4]), cur, with_customer=template != "receipt")
        render(template, tmp, inv)
        img = degrade(pdf_to_images(tmp, 200)[0], level, R) if level != "clean" else pdf_to_images(tmp, 200)[0]
        f = out / "images" / f"{template}_{level}_{i + 1:02d}.{ext}"
        if ext == "jpg":
            img.convert("RGB").save(f, "JPEG", quality=R.choice([55, 70, 85]))
        elif ext == "tiff":
            img.save(f, "TIFF", compression="tiff_deflate")
        else:
            img.save(f, "PNG")
        add(f, "image", [inv], tags=["image", ext, level, template])
    # multi-page TIFF (2-page invoice)
    inv = make_invoice(32, "GBP")
    render("classic", tmp, inv)
    pages = [degrade(p, "scan", R) for p in pdf_to_images(tmp, 200)]
    f = out / "images" / "classic_multipage.tiff"
    pages[0].save(f, "TIFF", save_all=True, append_images=pages[1:], compression="tiff_deflate")
    add(f, "image", [inv], tags=["image", "tiff", "multipage"])

    # 5) spreadsheets -----------------------------------------------------------------------
    reg_q1 = [make_invoice(R.randint(1, 3)) for _ in range(12)]
    reg_q2 = [make_invoice(R.randint(1, 3)) for _ in range(10)]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoices Q1"
    ws.append(["Accounts Payable Register"])
    ws.append(["Generated from finance system", None, None, "Confidential"])
    ws.append([])
    ws.append(["Invoice No", "Invoice Date", "Supplier", "Customer", "Net Amount", "VAT", "Total", "Currency"])
    for r in reg_q1:
        ws.append([r["invoice_number"], r["invoice_date"], r["supplier_name"], r["customer_name"], float(r["subtotal"]),
                   float(r["tax"]) if r["tax"] is not None else None, float(r["total"]), r["currency"]])
    ws2 = wb.create_sheet("Q2")
    ws2.append(["Bill No", "Date", "Vendor Name", "Client", "Subtotal", "Tax Amount", "Grand Total", "CCY"])
    for r in reg_q2:
        ws2.append([r["invoice_number"], r["invoice_date"], r["supplier_name"], r["customer_name"], float(r["subtotal"]),
                    float(r["tax"]) if r["tax"] is not None else None, float(r["total"]), r["currency"]])
    wb.create_sheet("Notes").append(["Reviewed by finance team"])
    f = out / "spreadsheets" / "ap_register_2025_multisheet.xlsx"
    wb.save(f)
    for r in reg_q1 + reg_q2:
        r["line_items"] = []
    add(f, "xlsx", reg_q1 + reg_q2, tags=["xlsx", "register", "multisheet", "title_rows"])

    lines_invs = [make_invoice(R.randint(2, 5)) for _ in range(8)]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice lines"
    ws.append(["Invoice #", "Date", "Supplier", "Customer", "Description", "Qty", "Unit Price", "Line Total",
               "Invoice Total", "Currency"])
    for inv in lines_invs:
        for j, li in enumerate(inv["line_items"]):
            ws.append([inv["invoice_number"], inv["invoice_date"], inv["supplier_name"], inv["customer_name"],
                       li["description"], float(li["quantity"]), float(li["unit_price"]), float(li["amount"]),
                       float(inv["total"]) if j == 0 else None, inv["currency"]])
    f = out / "spreadsheets" / "invoice_lines_export.xlsx"
    wb.save(f)
    for inv in lines_invs:
        inv["subtotal"] = inv["tax"] = None  # not present in the export
    add(f, "xlsx", lines_invs, tags=["xlsx", "line_items"])

    xls_invs = [make_invoice(R.randint(1, 3)) for _ in range(10)]
    book = xlwt.Workbook()
    sh = book.add_sheet("Register")
    date_style = xlwt.easyxf(num_format_str="DD/MM/YYYY")
    for c, h in enumerate(["Inv No", "Bill Date", "Vendor", "Customer Name", "Sub Total", "Tax", "Invoice Total",
                           "Currency"]):
        sh.write(0, c, h)
    for r_i, r in enumerate(xls_invs, start=1):
        sh.write(r_i, 0, r["invoice_number"])
        sh.write(r_i, 1, r["invoice_date"], date_style)
        sh.write(r_i, 2, r["supplier_name"])
        sh.write(r_i, 3, r["customer_name"])
        sh.write(r_i, 4, float(r["subtotal"]))
        if r["tax"] is not None:
            sh.write(r_i, 5, float(r["tax"]))
        sh.write(r_i, 6, float(r["total"]))
        sh.write(r_i, 7, r["currency"])
    f = out / "spreadsheets" / "legacy_register.xls"
    book.save(str(f))
    for r in xls_invs:
        r["line_items"] = []
    add(f, "xls", xls_invs, tags=["xls", "register"])

    eu_invs = [make_invoice(R.randint(1, 3), "EUR") for _ in range(8)]
    eu_invs[0]["supplier_name"] = "Müller Bürobedarf GmbH"
    eu_invs[3]["supplier_name"] = "Société Générale d'Équipement SA"
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["Rechnungsnummer / Invoice No", "Datum / Date", "Lieferant / Supplier", "Kunde / Customer",
                "Netto / Subtotal", "MwSt / VAT", "Gesamt / Total", "Währung / Currency"])
    for r in eu_invs:
        w.writerow([r["invoice_number"], r["invoice_date"].strftime("%d.%m.%Y"), r["supplier_name"], r["customer_name"],
                    fmt_amount(r["subtotal"], True), fmt_amount(r["tax"], True), fmt_amount(r["total"], True), "EUR"])
    f = out / "spreadsheets" / "eu_invoices_semicolon_latin1.csv"
    f.write_bytes(buf.getvalue().encode("latin-1"))
    for r in eu_invs:
        r["line_items"] = []
    add(f, "csv", eu_invs, tags=["csv", "semicolon", "latin1", "decimal_comma"])

    us_invs = [make_invoice(R.randint(1, 3), "USD") for _ in range(10)]
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["invoice_number", "invoice_date", "vendor", "customer", "subtotal", "tax", "total", "currency"])
    for r in us_invs:
        w.writerow([r["invoice_number"], r["invoice_date"].isoformat(), r["supplier_name"], r["customer_name"],
                    f"{r['subtotal']:.2f}", f"{r['tax']:.2f}" if r["tax"] is not None else "", f"{r['total']:.2f}", "USD"])
    w.writerow(["broken", "row", "with", "too", "many", "columns", "x", "y", "z", "extra"])
    f = out / "spreadsheets" / "us_invoices_utf8.csv"
    f.write_bytes(("﻿" + buf.getvalue()).encode("utf-8"))
    for r in us_invs:
        r["line_items"] = []
    add(f, "csv", us_invs, tags=["csv", "utf8_bom", "malformed_row"])

    # 6) validation scenarios ---------------------------------------------------------------
    inv = make_invoice(3, "USD")
    inv["total"] = q(inv["total"] + Decimal("100.00"))  # printed total is wrong
    f = out / "validation" / "total_mismatch.pdf"
    render("classic", f, inv)
    add(f, "pdf", [inv], issues=["TOTAL_MISMATCH"], tags=["validation"])

    inv = make_invoice(4, "GBP")
    inv["line_items"][1]["amount"] = q(inv["line_items"][1]["amount"] + Decimal("40.00"))  # line sum != subtotal
    f = out / "validation" / "line_items_mismatch.pdf"
    render("modern", f, inv)
    add(f, "pdf", [inv], issues=["LINE_ITEMS_MISMATCH", "LINE_ITEM_MATH"], tags=["validation"])

    inv = make_invoice(2, "USD")
    inv["invoice_number"] = None
    f = out / "validation" / "missing_invoice_number.pdf"
    render("classic", f, inv)
    add(f, "pdf", [inv], issues=["MISSING_FIELD"], tags=["validation", "missing_number"])

    dup = make_invoice(3, "EUR", number="DUP-55120", supplier="Brightstar Electronics GmbH")
    f1 = out / "validation" / "duplicate_a.pdf"
    render("european", f1, dup)
    add(f1, "pdf", [dup], issues=["DUPLICATE_INVOICE?"], tags=["validation", "duplicate_pair"])
    f2 = out / "validation" / "duplicate_b_rescan.png"
    render("european", tmp, dup)
    degrade(pdf_to_images(tmp, 200)[0], "scan", R).save(f2, "PNG")
    add(f2, "image", [dup], issues=["DUPLICATE_INVOICE?"], tags=["validation", "duplicate_pair"])

    # 7) corrupted / unsupported --------------------------------------------------------------
    good = (out / "pdf" / "classic_01.pdf").read_bytes()
    (out / "corrupted" / "truncated.pdf").write_bytes(good[: len(good) // 3])
    add(out / "corrupted" / "truncated.pdf", "pdf", [], expect="fail_or_degraded", error_code="CORRUPTED_FILE",
        tags=["corrupted"])
    (out / "corrupted" / "random_bytes.pdf").write_bytes(bytes(R.getrandbits(8) for _ in range(4096)))
    add(out / "corrupted" / "random_bytes.pdf", "pdf", [], expect="fail", error_code="CORRUPTED_FILE", tags=["corrupted"])
    (out / "corrupted" / "not_a_workbook.xlsx").write_bytes(b"PK\x03\x04" + bytes(R.getrandbits(8) for _ in range(2048)))
    add(out / "corrupted" / "not_a_workbook.xlsx", "xlsx", [], expect="fail", error_code="CORRUPTED_FILE",
        tags=["corrupted"])
    (out / "corrupted" / "empty.png").write_bytes(b"")
    add(out / "corrupted" / "empty.png", "image", [], expect="reject", error_code="EMPTY_DOCUMENT", tags=["corrupted"])
    (out / "corrupted" / "notes.txt").write_text("not an invoice")
    add(out / "corrupted" / "notes.txt", "other", [], expect="reject", error_code="UNSUPPORTED_FILE_TYPE",
        tags=["unsupported"])
    jpg = next(iter(sorted((out / "images").glob("*clean*.jpg"))), None)
    if jpg is not None:
        data = jpg.read_bytes()
        (out / "corrupted" / "truncated_photo.jpg").write_bytes(data[: int(len(data) * 0.8)])
        add(out / "corrupted" / "truncated_photo.jpg", "image", [], expect="fail_or_degraded", tags=["corrupted"])
    inv = make_invoice(2, "USD")
    render("classic", out / "corrupted" / "password_protected.pdf", inv, encrypt="s3cret")
    add(out / "corrupted" / "password_protected.pdf", "pdf", [], expect="fail", error_code="ENCRYPTED_FILE",
        tags=["encrypted"])

    tmp.unlink(missing_ok=True)

    # 8) stress files (not committed) ----------------------------------------------------------
    if args.stress:
        stress = out / "stress"
        stress.mkdir(exist_ok=True)
        big = fitz.open()
        stress_invs = []
        while big.page_count < 1000:
            template = ["classic", "modern", "statement", "european"][len(stress_invs) % 4]
            cur = TEMPLATE_CURRENCIES[template][len(stress_invs) % len(TEMPLATE_CURRENCIES[template])]
            inv = make_invoice(R.choice([3, 6, 10]), cur, number=f"STR-{100000 + len(stress_invs)}")
            render(template, tmp, inv)
            big.insert_pdf(fitz.open(tmp))
            stress_invs.append(inv)
        if big.page_count > 1000:
            big.delete_pages(1000, big.page_count - 1)
        big.save(stress / "stress_1000_pages_digital.pdf")
        imgs: list = []
        n_scanned_invoices = 0
        while len(imgs) < 100:
            inv = make_invoice(R.choice([3, 30]), "USD")
            render("classic", tmp, inv)
            imgs.extend(degrade(p, "scan", R) for p in pdf_to_images(tmp, 150))
            n_scanned_invoices += 1
        images_to_pdf(imgs[:100], stress / "stress_100_pages_scanned.pdf", 150)
        tmp.unlink(missing_ok=True)
        (stress / "stress_meta.json").write_text(json.dumps({"digital_pages": 1000, "digital_invoices_approx": len(stress_invs),
                                                               "scanned_pages": 100, "scanned_invoices_approx": n_scanned_invoices}, indent=2))

    (out / "labels.json").write_text(json.dumps({
        "generated": "2026-10-27", "seed": 20261027, "documents": labels, "reconciliation": expected_recon,
    }, indent=2, ensure_ascii=False))
    n_inv = sum(len(d["invoices"]) for d in labels)
    print(f"wrote {len(labels)} documents / {n_inv} labeled invoices to {out}")


if __name__ == "__main__":
    main()
